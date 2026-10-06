from __future__ import annotations

import hashlib
import io
import json
import re
import unicodedata
import uuid
import zipfile
from collections import Counter
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path, PurePosixPath
from typing import Any, Callable, Protocol

from .retrieval import embed_text, search_terms


@dataclass(frozen=True)
class CourseSpec:
    folder: str
    course_id: str
    course_name: str


COURSE_FOLDERS = (
    CourseSpec("Introduction To Artificial Intelligence", "CAI3101", "Introduction to Artificial Intelligence"),
    CourseSpec("Systems Programming", "CCS3202", "Systems Programming"),
    CourseSpec("Operating System", "CCS3203", "Operating Systems"),
    CourseSpec("Theory Of Computation", "CCS3402", "Theory of Computation"),
    CourseSpec("Prof. Training In Mobile Apps Programming", "CIT3200", "Professional Training in Mobile Apps Programming"),
    CourseSpec("Differential Equations", "EBA3202", "Differential Equations"),
)

SUPPORTED_SUFFIXES = {".pdf", ".pptx", ".docx", ".png", ".jpg", ".jpeg"}
PRIVATE_NAME_PATTERNS = (re.compile(r"\barwa\s+ahmed\b", re.IGNORECASE),)
_WEEK_RE = re.compile(r"(?:^|[^a-z])(?:week|w|lec(?:ture)?)\s*[-_,]?\s*(\d{1,2})(?:\b|_)", re.IGNORECASE)
_SAFE_NAME_RE = re.compile(r"[^A-Za-z0-9._()&+ -]+")


class MaterialSink(Protocol):
    def upsert_material_document(self, document: dict[str, Any]) -> None: ...
    def replace_material_chunks(self, document_id: str, chunks: list[dict[str, Any]]) -> None: ...
    def upsert_material_summary(self, document_id: str, summary: dict[str, Any]) -> None: ...
    def save_material_ingestion_job(self, job: dict[str, Any]) -> None: ...


@dataclass
class ExtractedSection:
    locator: int
    text: str


def _normalized_text(value: str) -> str:
    value = unicodedata.normalize("NFKC", value).replace("\x00", " ")
    return re.sub(r"\s+", " ", value).strip()


def _safe_filename(value: str) -> str:
    normalized = _SAFE_NAME_RE.sub("-", _normalized_text(value)).strip(" .-")
    return normalized[:140] or "material"


def _course_for_member(member: PurePosixPath) -> CourseSpec | None:
    parts = [part.casefold() for part in member.parts]
    for course in COURSE_FOLDERS:
        if course.folder.casefold() in parts:
            return course
    return None


def _should_exclude(name: str) -> str | None:
    normalized = unicodedata.normalize("NFKC", name)
    lowered = normalized.casefold()
    if "__macosx/" in lowered or lowered.endswith(".ds_store"):
        return "system metadata"
    if any(pattern.search(normalized) for pattern in PRIVATE_NAME_PATTERNS):
        return "possible student submission"
    return None


def _category(member: PurePosixPath) -> str:
    path = "/".join(member.parts).casefold()
    name = member.name.casefold()
    if any(token in path for token in ("exam", "testbank")) or any(
        token in name for token in ("quiz", "model answer", "old ai")
    ):
        return "exam_practice"
    if "project" in path:
        return "project"
    if "lab" in path:
        return "lab"
    if any(token in path for token in ("/sections/", "/sec/", "/sheets/")):
        return "section"
    if "assignment" in name:
        return "assignment"
    if any(token in path for token in ("/lectures/", "/lec/", "/ipad/")) or re.match(
        r"^(?:w|week|lec)", name
    ):
        return "lecture"
    return "reference"


def _week_number(name: str) -> int | None:
    match = _WEEK_RE.search(name)
    if not match:
        return None
    week = int(match.group(1))
    return week if 1 <= week <= 16 else None


def _extract_pdf(data: bytes) -> tuple[list[ExtractedSection], int]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data), strict=False)
    sections = []
    for index, page in enumerate(reader.pages, start=1):
        text = _normalized_text(page.extract_text() or "")
        if text:
            sections.append(ExtractedSection(index, text))
    return sections, len(reader.pages)


def _extract_pptx(data: bytes) -> tuple[list[ExtractedSection], int]:
    from pptx import Presentation

    deck = Presentation(io.BytesIO(data))
    sections: list[ExtractedSection] = []
    for index, slide in enumerate(deck.slides, start=1):
        values: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                values.extend(paragraph.text for paragraph in shape.text_frame.paragraphs)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    values.append(" | ".join(cell.text for cell in row.cells))
        text = _normalized_text("\n".join(values))
        if text:
            sections.append(ExtractedSection(index, text))
    return sections, len(deck.slides)


def _extract_docx(data: bytes) -> tuple[list[ExtractedSection], int]:
    from docx import Document

    document = Document(io.BytesIO(data))
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            values.append(" | ".join(cell.text for cell in row.cells))
    text = _normalized_text("\n".join(values))
    return ([ExtractedSection(1, text)] if text else []), 1


def extract_sections(data: bytes, suffix: str) -> tuple[list[ExtractedSection], int]:
    if suffix == ".pdf":
        return _extract_pdf(data)
    if suffix == ".pptx":
        return _extract_pptx(data)
    if suffix == ".docx":
        return _extract_docx(data)
    return [], 1


def _split_long_text(text: str, maximum: int = 3_200) -> list[str]:
    if len(text) <= maximum:
        return [text]
    sentences = re.split(r"(?<=[.!?])\s+", text)
    pieces: list[str] = []
    current = ""
    for sentence in sentences:
        if current and len(current) + len(sentence) + 1 > maximum:
            pieces.append(current)
            current = current[-240:] + " " + sentence
        else:
            current = f"{current} {sentence}".strip()
    if current:
        pieces.append(current)
    return pieces


def build_chunks(
    document_id: str,
    course: CourseSpec,
    title: str,
    category: str,
    sections: list[ExtractedSection],
) -> list[dict[str, Any]]:
    chunks: list[dict[str, Any]] = []
    for section in sections:
        for text in _split_long_text(section.text):
            content = _normalized_text(text)
            if len(content) < 40:
                continue
            chunk_index = len(chunks)
            embedding_input = f"{course.course_name}\n{title}\n{category}\n{content}"
            chunks.append(
                {
                    "chunk_id": hashlib.sha256(
                        f"{document_id}:{chunk_index}".encode("utf-8")
                    ).hexdigest(),
                    "document_id": document_id,
                    "chunk_index": chunk_index,
                    "content": content,
                    "page_start": section.locator,
                    "page_end": section.locator,
                    "token_count": max(1, len(content) // 4),
                    "embedding": embed_text(embedding_input),
                }
            )
    return chunks


def summarize_chunks(chunks: list[dict[str, Any]]) -> dict[str, Any]:
    combined = " ".join(chunk["content"] for chunk in chunks[:6])
    terms = [term for term in search_terms(combined) if not term.isdigit()]
    stop = {
        "about", "after", "also", "been", "before", "between", "course",
        "from", "have", "into", "more", "that", "their", "these", "this",
        "through", "using", "were", "what", "when", "where", "which", "with",
    }
    keywords = [term for term, _ in Counter(term for term in terms if term not in stop).most_common(12)]
    return {
        "summary": combined[:1_200].strip(),
        "learning_objectives": [],
        "keywords": keywords,
        "generator": "deterministic-extractive-v1",
    }


def ingest_semester_archive(
    archive_path: Path,
    storage_root: Path,
    sink: MaterialSink,
    progress: Callable[[str], None] | None = None,
) -> dict[str, Any]:
    """Ingest the approved CS Semester 5 archive into a repository sink."""
    job_id = str(uuid.uuid4())
    started_at = datetime.now(UTC).isoformat()
    report: dict[str, Any] = {
        "job_id": job_id,
        "source": str(archive_path),
        "started_at": started_at,
        "status": "running",
        "documents": 0,
        "chunks": 0,
        "duplicates": 0,
        "excluded": [],
        "failed": [],
        "by_course": {},
    }
    sink.save_material_ingestion_job(report)
    storage_root.mkdir(parents=True, exist_ok=True)
    seen_checksums: set[tuple[str, str]] = set()

    try:
        with zipfile.ZipFile(archive_path) as archive:
            members = sorted(
                (item for item in archive.infolist() if not item.is_dir()),
                key=lambda item: item.filename.casefold(),
            )
            for item in members:
                member = PurePosixPath(item.filename)
                reason = _should_exclude(item.filename)
                course = _course_for_member(member)
                suffix = member.suffix.casefold()
                if reason or course is None or suffix not in SUPPORTED_SUFFIXES:
                    report["excluded"].append(
                        {"path": item.filename, "reason": reason or "unsupported or outside a course"}
                    )
                    continue

                try:
                    data = archive.read(item)
                    checksum = hashlib.sha256(data).hexdigest()
                    duplicate_key = (course.course_id, checksum)
                    if duplicate_key in seen_checksums:
                        report["duplicates"] += 1
                        continue
                    seen_checksums.add(duplicate_key)

                    category = _category(member)
                    week = _week_number(member.name)
                    title = _normalized_text(member.stem)
                    document_id = hashlib.sha256(
                        f"{course.course_id}:{item.filename}".encode("utf-8")
                    ).hexdigest()
                    relative_path = Path(
                        "cs",
                        "semester-5",
                        course.course_id,
                        category,
                        f"{checksum[:12]}-{_safe_filename(member.name)}",
                    )
                    target = storage_root / relative_path
                    target.parent.mkdir(parents=True, exist_ok=True)
                    if not target.exists() or target.stat().st_size != len(data):
                        with target.open("wb") as stream:
                            stream.write(data)

                    sections, page_count = extract_sections(data, suffix)
                    chunks = build_chunks(document_id, course, title, category, sections)
                    now = datetime.now(UTC).isoformat()
                    document_record = {
                        "document_id": document_id,
                        "course_id": course.course_id,
                        "course_name": course.course_name,
                        "major_code": "CS",
                        "program_semester": 5,
                        "title": title,
                        "category": category,
                        "week_number": week,
                        "original_filename": member.name,
                        "source_archive_path": item.filename,
                        "storage_provider": "local",
                        "storage_path": relative_path.as_posix(),
                        "mime_type": {
                            ".pdf": "application/pdf",
                            ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                            ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            ".png": "image/png",
                            ".jpg": "image/jpeg",
                            ".jpeg": "image/jpeg",
                        }[suffix],
                        "checksum": checksum,
                        "size_bytes": len(data),
                        "page_count": page_count,
                        "visibility": "student_practice" if category == "exam_practice" else "student",
                        "status": "ready" if chunks else "needs_ocr",
                        "created_at": now,
                        "updated_at": now,
                    }
                    sink.upsert_material_document(document_record)
                    sink.replace_material_chunks(document_id, chunks)
                    sink.upsert_material_summary(document_id, summarize_chunks(chunks))
                    report["documents"] += 1
                    report["chunks"] += len(chunks)
                    course_stats = report["by_course"].setdefault(
                        course.course_id, {"documents": 0, "chunks": 0}
                    )
                    course_stats["documents"] += 1
                    course_stats["chunks"] += len(chunks)
                    if progress and report["documents"] % 10 == 0:
                        progress(
                            f"Imported {report['documents']} documents and "
                            f"{report['chunks']} searchable sections"
                        )
                except Exception as error:
                    report["failed"].append(
                        {"path": item.filename, "error": f"{type(error).__name__}: {error}"}
                    )
    except Exception:
        report["status"] = "failed"
        report["finished_at"] = datetime.now(UTC).isoformat()
        sink.save_material_ingestion_job(report)
        raise

    report["status"] = "completed_with_errors" if report["failed"] else "completed"
    report["finished_at"] = datetime.now(UTC).isoformat()
    sink.save_material_ingestion_job(report)
    report_path = storage_root / "cs" / "semester-5" / "ingestion-report.json"
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    return report
